from pathlib import Path
from shutil import copy2

import openpyxl

from django.core.management.base import BaseCommand

from cargadores.models import CargadorElectrico


class Command(BaseCommand):

    help = (
        "Actualiza latitud y longitud del Excel "
        "usando los cargadores de la base de datos local"
    )

    def handle(self, *args, **options):

        ruta_excel = (
            Path(__file__)
            .resolve()
            .parent
            .parent
            .parent
            / "data"
            / "cargadores_electricos_chile_completo.xlsx"
        )


        if not ruta_excel.exists():

            self.stdout.write(
                self.style.ERROR(
                    f"No se encontró el archivo: {ruta_excel}"
                )
            )

            return


        # =========================================
        # CREAR RESPALDO DEL EXCEL
        # =========================================

        ruta_respaldo = (
            ruta_excel.parent
            / "cargadores_electricos_chile_completo_backup.xlsx"
        )


        copy2(
            ruta_excel,
            ruta_respaldo
        )


        self.stdout.write(
            f"Respaldo creado: {ruta_respaldo}"
        )


        # =========================================
        # ABRIR EXCEL
        # =========================================

        libro = openpyxl.load_workbook(
            ruta_excel
        )


        hoja = libro[
            "Puntos_recarga_Chile"
        ]


        actualizados = 0
        sin_coordenadas = 0
        no_encontrados = 0


        # =========================================
        # RECORRER FILAS DEL EXCEL
        # =========================================

        for fila in range(
            2,
            hoja.max_row + 1
        ):

            nombre = hoja.cell(
                row=fila,
                column=6
            ).value

            direccion = hoja.cell(
                row=fila,
                column=7
            ).value


            if not nombre:

                continue


            try:

                cargador = (
                    CargadorElectrico.objects.get(
                        nombre=nombre,
                        direccion=direccion
                    )
                )

            except CargadorElectrico.DoesNotExist:

                no_encontrados += 1

                continue


            if (
                cargador.latitud is None
                or
                cargador.longitud is None
            ):

                sin_coordenadas += 1

                continue


            # Columna K = Latitud
            hoja.cell(
                row=fila,
                column=11
            ).value = float(
                cargador.latitud
            )


            # Columna L = Longitud
            hoja.cell(
                row=fila,
                column=12
            ).value = float(
                cargador.longitud
            )


            actualizados += 1


        # =========================================
        # GUARDAR EXCEL
        # =========================================

        libro.save(
            ruta_excel
        )


        self.stdout.write(
            self.style.SUCCESS(
                f"Filas actualizadas: {actualizados}"
            )
        )


        self.stdout.write(
            self.style.WARNING(
                f"Filas sin coordenadas: {sin_coordenadas}"
            )
        )


        self.stdout.write(
            self.style.WARNING(
                f"Registros no encontrados: {no_encontrados}"
            )
        )


        self.stdout.write(
            self.style.SUCCESS(
                "Excel actualizado correctamente."
            )
        )
        